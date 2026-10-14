"""Content-addressed ingestion into normalized Parquet snapshots."""

from __future__ import annotations

import hashlib
import math
import os
import shutil
import uuid
from collections.abc import Callable, Iterable, Iterator
from datetime import UTC, datetime
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook  # type: ignore[import-untyped]

from ai_powerbi_analyst.contracts import (
    ColumnDataType,
    SourceFile,
    SourceKind,
    SourceManifest,
    TableProfile,
)
from ai_powerbi_analyst.data.onboarding import (
    _detect_delimiter,
    _identifier,
    _profile_frames,
    _safe_filename,
)
from ai_powerbi_analyst.integrity import sha256_file

from .contracts import (
    DataQualityFinding,
    DatasetSnapshot,
    NormalizedTableSnapshot,
    SourceInventoryItem,
)
from .inventory import csv_encoding, detect_format


class StudioIngestionError(RuntimeError):
    """Selected files could not be normalized safely."""


class StudioIngestionService:
    def __init__(self, state_root: str | Path) -> None:
        self.state_root = Path(state_root).expanduser().resolve(strict=False)
        self.originals_root = self.state_root / "sources" / "original"
        self.snapshots_root = self.state_root / "snapshots"
        self.manifests_root = self.state_root / "manifests"

    def ingest(
        self,
        workspace_id: str,
        items: Iterable[SourceInventoryItem],
        *,
        table_names: dict[str, str] | None = None,
        column_types: dict[str, dict[str, ColumnDataType]] | None = None,
        column_type_warnings: dict[str, dict[str, str]] | None = None,
        utc_columns: dict[str, set[str]] | None = None,
    ) -> tuple[DatasetSnapshot, SourceManifest]:
        selected = list(items)
        if not selected:
            raise StudioIngestionError("Select at least one supported local data file")
        required = sum(item.size_bytes for item in selected) * 2 + 256 * 1024**2
        free = shutil.disk_usage(self.state_root.parent).free
        if free < required:
            raise StudioIngestionError(
                "Not enough free disk for immutable and normalized snapshots"
            )
        snapshot_id = f"snapshot-{uuid.uuid4().hex}"
        normalized_root = self.snapshots_root / snapshot_id
        normalized_root.mkdir(parents=True, exist_ok=False)
        normalized_tables: list[NormalizedTableSnapshot] = []
        findings: list[DataQualityFinding] = []
        sources: list[SourceFile] = []
        profiles: list[TableProfile] = []
        try:
            for item in selected:
                path = Path(item.absolute_path).resolve(strict=True)
                detected, warning = detect_format(path)
                if detected not in {"csv", "xls", "xlsx", "xlsm", "parquet"}:
                    raise StudioIngestionError(f"Unsupported source signature: {path.name}")
                original_sha = sha256_file(path)
                original_kind = {
                    "csv": SourceKind.CSV,
                    "xls": SourceKind.EXCEL,
                    "xlsx": SourceKind.EXCEL,
                    "xlsm": SourceKind.EXCEL,
                    "parquet": SourceKind.PARQUET,
                }[detected]
                original_encoding = csv_encoding(path) if detected == "csv" else None
                original_delimiter = (
                    _detect_delimiter(path, original_encoding)
                    if detected == "csv" and original_encoding is not None
                    else None
                )
                original_root = self.originals_root / original_sha
                original_root.mkdir(parents=True, exist_ok=True)
                original_path = original_root / _safe_filename(path.name)
                if not original_path.exists():
                    temporary = original_path.with_suffix(original_path.suffix + ".tmp")
                    shutil.copyfile(path, temporary)
                    if sha256_file(temporary) != original_sha:
                        temporary.unlink(missing_ok=True)
                        raise StudioIngestionError(
                            "Original source snapshot hash verification failed"
                        )
                    os.replace(temporary, original_path)
                elif sha256_file(original_path) != original_sha:
                    raise StudioIngestionError("Preserved original source hash verification failed")
                # Both passes must read the reviewed bytes, even if the user or
                # another application changes the live source during profiling.
                frames = self._frames(original_path, detected)
                for index, (table_name, sheet, frame_factory) in enumerate(frames, start=1):
                    if detected in {"csv", "parquet"}:
                        table_name = path.stem
                    table_name = (table_names or {}).get(str(path), table_name)
                    stable = _identifier(
                        f"{Path(path.name).stem}_{table_name}_{index}_{item.id[-8:]}"
                    )
                    profile = _profile_frames(
                        frame_factory(),
                        stable,
                        table_name,
                        sheet=sheet,
                        source_id=stable,
                        declared_types=(column_types or {}).get(str(path)),
                    )
                    declared_names = list((column_types or {}).get(str(path), {}))
                    source_names = (
                        declared_names
                        if len(declared_names) == len(profile.columns)
                        else [column.name for column in profile.columns]
                    )
                    normalized_utc_columns: set[str] = set()
                    for source_name, column in zip(source_names, profile.columns, strict=True):
                        note = (column_type_warnings or {}).get(str(path), {}).get(source_name)
                        if note:
                            findings.append(
                                _finding("mixed_types", "info", profile.id, column.id, note)
                            )
                        if source_name in (utc_columns or {}).get(str(path), set()):
                            normalized_utc_columns.add(column.name)
                    if profile.row_count == 0:
                        findings.append(
                            _finding(
                                "empty_table",
                                "warning",
                                profile.id,
                                None,
                                f"{table_name} has no rows",
                            )
                        )
                    parquet_path = normalized_root / f"{profile.id}.parquet"
                    _write_normalized_parquet(
                        frame_factory(), profile, parquet_path, utc_columns=normalized_utc_columns
                    )
                    normalized_sha = sha256_file(parquet_path)
                    source_id = profile.source_id
                    source = SourceFile(
                        id=source_id,
                        original_name=f"{path.name} :: {table_name}",
                        original_path=str(path),
                        snapshot_path=str(parquet_path),
                        sha256=normalized_sha,
                        original_sha256=original_sha,
                        size_bytes=parquet_path.stat().st_size,
                        kind=SourceKind.PARQUET,
                        original_kind=original_kind,
                        sheets=[sheet] if sheet else [],
                        encoding=original_encoding,
                        delimiter=original_delimiter,
                    )
                    sources.append(source)
                    profiles.append(profile)
                    normalized_tables.append(
                        NormalizedTableSnapshot(
                            id=profile.id,
                            source_item_id=item.id,
                            original_sha256=original_sha,
                            original_snapshot_path=str(original_path),
                            normalized_sha256=normalized_sha,
                            normalized_parquet_path=str(parquet_path),
                            table=profile,
                        )
                    )
                if warning:
                    findings.append(_finding("mixed_types", "info", None, None, warning))
                if detected == "xlsm":
                    findings.append(
                        _finding(
                            "macros_ignored",
                            "info",
                            None,
                            None,
                            (
                                f"Macros in {path.name} were preserved as original bytes "
                                "but never executed"
                            ),
                        )
                    )
            if not profiles:
                raise StudioIngestionError("Selected sources contain no readable tables")
            created_at = datetime.now(UTC)
            manifest = SourceManifest(
                import_id=snapshot_id,
                imported_at=created_at,
                sources=sources,
                tables=profiles,
            )
            self.manifests_root.mkdir(parents=True, exist_ok=True)
            manifest_path = self.manifests_root / f"{snapshot_id}.json"
            manifest_path.write_text(manifest.model_dump_json(indent=2) + "\n", encoding="utf-8")
            snapshot_hash = hashlib.sha256(
                "".join(sorted(item.normalized_sha256 for item in normalized_tables)).encode(
                    "ascii"
                )
            ).hexdigest()
            snapshot = DatasetSnapshot(
                id=snapshot_id,
                workspace_id=workspace_id,
                created_at=created_at,
                original_bytes=sum(item.size_bytes for item in selected),
                normalized_bytes=sum(
                    Path(item.normalized_parquet_path).stat().st_size for item in normalized_tables
                ),
                tables=normalized_tables,
                findings=findings,
                manifest_path=str(manifest_path),
                base_hash=snapshot_hash,
            )
            return snapshot, manifest
        except Exception:
            shutil.rmtree(normalized_root, ignore_errors=True)
            raise

    def load_manifest(self, snapshot: DatasetSnapshot) -> SourceManifest:
        return SourceManifest.model_validate_json(
            Path(snapshot.manifest_path).read_text(encoding="utf-8")
        )

    def relink_manifest(self, snapshot: DatasetSnapshot) -> SourceManifest:
        """Return provenance-complete bindings for reviewed original local files."""
        manifest = self.load_manifest(snapshot)
        snapshot_by_source = {item.table.source_id: item for item in snapshot.tables}
        enriched: list[SourceFile] = []
        for source in manifest.sources:
            if not source.original_path:
                raise StudioIngestionError(
                    f"Source has no reviewed original path: {source.original_name}"
                )
            original = Path(source.original_path).expanduser().resolve(strict=True)
            detected, _warning = detect_format(original)
            if detected not in {"csv", "xls", "xlsx", "xlsm", "parquet"}:
                raise StudioIngestionError(
                    f"Original source format is unsupported: {source.original_name}"
                )
            original_kind = {
                "csv": SourceKind.CSV,
                "xls": SourceKind.EXCEL,
                "xlsx": SourceKind.EXCEL,
                "xlsm": SourceKind.EXCEL,
                "parquet": SourceKind.PARQUET,
            }[detected]
            original_snapshot = snapshot_by_source.get(source.id)
            original_sha = source.original_sha256 or (
                original_snapshot.original_sha256 if original_snapshot is not None else None
            )
            if original_sha is None:
                raise StudioIngestionError(
                    f"Source has no reviewed original hash: {source.original_name}"
                )
            encoding = csv_encoding(original) if detected == "csv" else None
            delimiter = (
                _detect_delimiter(original, encoding)
                if detected == "csv" and encoding is not None
                else None
            )
            enriched.append(
                source.model_copy(
                    update={
                        "original_sha256": original_sha,
                        "original_kind": original_kind,
                        "encoding": encoding,
                        "delimiter": delimiter,
                    }
                )
            )
        return manifest.model_copy(update={"sources": enriched})

    def _frames(
        self,
        path: Path,
        detected: str,
    ) -> list[tuple[str, str | None, Callable[[], Iterable[pd.DataFrame]]]]:
        if detected == "csv":
            encoding = csv_encoding(path)
            delimiter = _detect_delimiter(path, encoding)

            def csv_frames() -> Iterable[pd.DataFrame]:
                try:
                    return pd.read_csv(
                        path,
                        encoding=encoding,
                        sep=delimiter,
                        chunksize=25_000,
                        dtype=object,
                        keep_default_na=True,
                    )
                except (OSError, UnicodeError, pd.errors.ParserError, ValueError) as error:
                    raise StudioIngestionError(
                        f"Could not read CSV {path.name}: {error}"
                    ) from error

            return [(path.stem, None, csv_frames)]
        if detected == "parquet":

            def parquet_frames() -> Iterable[pd.DataFrame]:
                parquet = pq.ParquetFile(path)
                if parquet.metadata.num_rows == 0:
                    return [
                        parquet.schema_arrow.empty_table().to_pandas(types_mapper=pd.ArrowDtype)
                    ]
                return (
                    batch.to_pandas(types_mapper=pd.ArrowDtype)
                    for batch in parquet.iter_batches(batch_size=25_000)
                )

            return [
                (
                    path.stem,
                    None,
                    parquet_frames,
                )
            ]
        if detected == "xls":
            try:
                workbook = pd.ExcelFile(path, engine="calamine")
                sheet_names = [str(sheet) for sheet in workbook.sheet_names]
                workbook.close()
                return [(sheet, sheet, _calamine_factory(path, sheet)) for sheet in sheet_names]
            except (OSError, ValueError, ImportError) as error:
                raise StudioIngestionError(
                    f"Could not read legacy Excel workbook {path.name}: {error}"
                ) from error
        return self._openpyxl_frames(path)

    def _openpyxl_frames(
        self,
        path: Path,
    ) -> list[tuple[str, str | None, Callable[[], Iterable[pd.DataFrame]]]]:
        try:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            sheet_names = [sheet.title for sheet in workbook.worksheets]
            workbook.close()
            return [(sheet, sheet, _openpyxl_factory(path, sheet)) for sheet in sheet_names]
        except (OSError, ValueError, ImportError) as error:
            raise StudioIngestionError(
                f"Could not read Excel workbook {path.name}: {error}"
            ) from error


def _write_normalized_parquet(
    frames: Iterable[pd.DataFrame],
    profile: TableProfile,
    path: Path,
    *,
    utc_columns: set[str] | None = None,
) -> None:
    for column in profile.columns:
        if column.data_type.value == "integer" and any(
            isinstance(bound, int) and not -(2**63) <= bound <= 2**63 - 1
            for bound in (column.minimum, column.maximum)
        ):
            raise StudioIngestionError(
                f"Column {column.name} in table {profile.name} contains an integer outside "
                "the supported signed 64-bit range. Store this column as text in the "
                "source file or database view, then import it again to preserve its values."
            )
    writer: pq.ParquetWriter | None = None
    try:
        for frame in frames:
            normalized = frame.copy()
            normalized.columns = [column.name for column in profile.columns]
            for column in profile.columns:
                series = normalized[column.name]
                if column.data_type.value == "integer":
                    normalized[column.name] = pd.to_numeric(
                        series, errors="coerce", dtype_backend="numpy_nullable"
                    ).astype("Int64")
                elif column.data_type.value == "number":
                    normalized[column.name] = pd.to_numeric(series, errors="coerce").astype(
                        "Float64"
                    )
                elif column.data_type.value in {"date", "datetime"}:
                    normalized[column.name] = pd.to_datetime(
                        series,
                        errors="coerce",
                        format="mixed",
                        utc=column.name in (utc_columns or set()),
                    )
                elif column.data_type.value == "boolean":
                    normalized[column.name] = series.map(_to_bool).astype("boolean")
                else:
                    normalized[column.name] = series.astype("string")
            table = pa.Table.from_pandas(normalized, preserve_index=False)
            if writer is None:
                writer = pq.ParquetWriter(path, table.schema, compression="zstd")
            writer.write_table(table)
        if writer is None:
            raise StudioIngestionError(f"Table {profile.name} has no data batches")
    finally:
        if writer is not None:
            writer.close()


def _calamine_factory(path: Path, sheet: str) -> Callable[[], Iterable[pd.DataFrame]]:
    def frames() -> Iterable[pd.DataFrame]:
        try:
            return (pd.read_excel(path, sheet_name=sheet, dtype=object, engine="calamine"),)
        except (OSError, ValueError, ImportError) as error:
            raise StudioIngestionError(
                f"Could not read legacy Excel sheet {sheet} from {path.name}: {error}"
            ) from error

    return frames


def _openpyxl_factory(path: Path, sheet: str) -> Callable[[], Iterable[pd.DataFrame]]:
    return lambda: _iter_openpyxl_sheet(path, sheet)


def _iter_openpyxl_sheet(path: Path, sheet_name: str) -> Iterator[pd.DataFrame]:
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return
        names = _unique_headers(header)
        buffer: list[tuple[object, ...]] = []
        produced = False
        for row in rows:
            buffer.append(tuple(row))
            if len(buffer) >= 25_000:
                yield pd.DataFrame(buffer, columns=names)
                buffer = []
                produced = True
        if buffer or not produced:
            yield pd.DataFrame(buffer, columns=names)
    finally:
        workbook.close()


def _to_bool(value: object) -> bool | None:
    if value is None or value is pd.NA or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {"true", "yes", "1"}


def _unique_headers(values: Iterable[object]) -> list[str]:
    used: set[str] = set()
    result: list[str] = []
    for index, value in enumerate(values, start=1):
        base = " ".join(str(value if value is not None else "").split()) or f"Column {index}"
        name, suffix = base, 2
        while name in used:
            name = f"{base} {suffix}"
            suffix += 1
        used.add(name)
        result.append(name)
    return result


def _finding(
    code: str,
    severity: str,
    table_id: str | None,
    column_id: str | None,
    message: str,
) -> DataQualityFinding:
    digest = hashlib.sha256(f"{code}:{table_id}:{column_id}:{message}".encode()).hexdigest()[:16]
    return DataQualityFinding(
        id=f"finding-{digest}",
        code=code,  # type: ignore[arg-type]
        severity=severity,  # type: ignore[arg-type]
        table_id=table_id,
        column_id=column_id,
        message=message,
    )
