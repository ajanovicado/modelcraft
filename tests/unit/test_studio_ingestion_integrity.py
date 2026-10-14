from __future__ import annotations

from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from ai_powerbi_analyst.contracts import ColumnDataType
from ai_powerbi_analyst.data.onboarding import _profile_frames
from ai_powerbi_analyst.studio import ingestion as module
from ai_powerbi_analyst.studio.ingestion import StudioIngestionError, StudioIngestionService
from ai_powerbi_analyst.studio.inventory import inventory_folder


def test_profile_and_normalization_read_same_preserved_bytes_when_live_file_changes(
    tmp_path, monkeypatch
):
    selected = tmp_path / "selected"
    selected.mkdir()
    source = selected / "ventas_α.csv"
    original = b"ID,Sales\n1,10\n2,20\n"
    source.write_bytes(original)
    item = inventory_folder(selected)[0]
    profile = module._profile_frames

    def change_live_after_profile(*args, **kwargs):
        result = profile(*args, **kwargs)
        source.write_text("ID,Sales\n1,999\n", encoding="utf-8")
        return result

    monkeypatch.setattr(module, "_profile_frames", change_live_after_profile)
    snapshot, manifest = StudioIngestionService(tmp_path / "state").ingest("workspace-qa", [item])
    table = snapshot.tables[0]
    assert table.table.row_count == 2
    assert table.table.name == "ventas_α"
    assert Path(table.original_snapshot_path).read_bytes() == original
    assert pq.read_table(table.normalized_parquet_path).to_pydict() == {
        "ID": [1, 2],
        "Sales": [10, 20],
    }
    assert manifest.sources[0].original_path == str(source)


def test_corrupt_cached_original_cannot_supply_a_new_snapshot(tmp_path):
    selected = tmp_path / "selected"
    selected.mkdir()
    source = selected / "data.csv"
    source.write_text("ID,Sales\n1,10\n", encoding="utf-8")
    item = inventory_folder(selected)[0]
    service = StudioIngestionService(tmp_path / "state")
    first, _ = service.ingest("workspace-qa", [item])
    Path(first.tables[0].original_snapshot_path).write_text("ID,Sales\n1,999\n", encoding="utf-8")
    with pytest.raises(StudioIngestionError, match="Preserved original source hash"):
        service.ingest("workspace-qa", [item])
    assert len(list(service.snapshots_root.iterdir())) == 1


def test_nullable_parquet_integer_is_not_converted_through_float(tmp_path):
    selected = tmp_path / "selected"
    selected.mkdir()
    source = selected / "ids.parquet"
    pq.write_table(pa.table({"ID": pa.array([123456789012345678, None], type=pa.int64())}), source)
    snapshot, _ = StudioIngestionService(tmp_path / "state").ingest(
        "workspace-qa", inventory_folder(selected)
    )
    table = snapshot.tables[0]
    assert table.table.columns[0].data_type is ColumnDataType.INTEGER
    assert table.table.columns[0].sample_values == [123456789012345678]
    assert pq.read_table(table.normalized_parquet_path).to_pydict() == {
        "ID": [123456789012345678, None]
    }


def test_unsigned_parquet_integer_outside_signed64_fails_without_wrapping(tmp_path):
    selected = tmp_path / "selected"
    selected.mkdir()
    source = selected / "ids.parquet"
    pq.write_table(pa.table({"ID": pa.array([2**64 - 1, None], type=pa.uint64())}), source)
    service = StudioIngestionService(tmp_path / "state")
    with pytest.raises(StudioIngestionError, match="signed 64-bit range.*as text"):
        service.ingest("workspace-qa", inventory_folder(selected))
    assert not list(service.snapshots_root.iterdir())
    assert pq.read_table(source).to_pydict() == {"ID": [2**64 - 1, None]}


@pytest.mark.parametrize("value", [-(2**63) - 1, 2**63])
def test_csv_integer_outside_signed64_fails_with_recovery_instruction(tmp_path, value):
    selected = tmp_path / "selected"
    selected.mkdir()
    (selected / "ids.csv").write_text(f"ID\n{value}\n", encoding="utf-8")
    service = StudioIngestionService(tmp_path / "state")
    with pytest.raises(StudioIngestionError, match="signed 64-bit range.*as text"):
        service.ingest("workspace-qa", inventory_folder(selected))
    assert not list(service.snapshots_root.iterdir())


def test_excel_headers_remain_unique_without_overwriting_cells(tmp_path):
    selected = tmp_path / "selected"
    selected.mkdir()
    source = selected / "duplicates.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "A", "A 2", 0, None, "Column 5"])
    sheet.append([10, 20, 30, 40, 50, 60])
    workbook.save(source)
    workbook.close()
    snapshot, _ = StudioIngestionService(tmp_path / "state").ingest(
        "workspace-qa", inventory_folder(selected)
    )
    values = pq.read_table(snapshot.tables[0].normalized_parquet_path).to_pydict()
    assert list(values) == ["A", "A 2", "A 2 2", "0", "Column 5", "Column 5 2"]
    assert list(values.values()) == [[10], [20], [30], [40], [50], [60]]


def test_normalized_header_collision_preserves_declared_source_types():
    frame = pd.DataFrame([["001", "002", "003"]], columns=["A", "A ", "A 2"])
    profile = _profile_frames(
        [frame],
        "ids",
        "IDs",
        sheet=None,
        declared_types={name: ColumnDataType.STRING for name in frame.columns},
    )
    assert [column.name for column in profile.columns] == ["A", "A 2", "A 2 2"]
    assert len({column.id for column in profile.columns}) == 3
    assert [column.sample_values for column in profile.columns] == [["001"], ["002"], ["003"]]
    assert all(column.data_type is ColumnDataType.STRING for column in profile.columns)
